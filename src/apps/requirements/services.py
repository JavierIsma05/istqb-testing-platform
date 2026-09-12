import csv
import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass

from .models import Requirement


class RequirementPdfImportError(Exception):
    pass


class RequirementSourceImportError(RequirementPdfImportError):
    pass


@dataclass
class ParsedRequirement:
    title: str
    description: str
    acceptance_criteria: str = ''
    requirement_type: str = Requirement.RequirementType.FUNCTIONAL
    priority: str = Requirement.Priority.MEDIUM
    status: str = Requirement.Status.PENDING


REQUIREMENT_START_RE = re.compile(
    r'^\s*(?:(?P<code>(?:REQ|RF|RNF|FR|NFR)[-\s]?\d+)\s*[:.)-]\s*)?(?P<body>.+)$',
    re.IGNORECASE,
)
NUMBERED_START_RE = re.compile(r'^\s*\d+[\).:-]\s+.+$')
REQUIREMENT_WORD_RE = re.compile(
    r'\b(el sistema|la aplicacion|la plataforma|se debe|debe|debera|permitir|permitira|permite|registrar|mostrar|generar|requisito)\b',
    re.IGNORECASE,
)
CODE_PREFIX_RE = re.compile(r'^(?P<prefix>REQ|RF|RNF|FR|NFR)[-\s]*$', re.IGNORECASE)
CODE_NUMBER_RE = re.compile(r'^\d{1,4}$')
COMPLETE_CODE_RE = re.compile(r'^(?P<prefix>REQ|RF|RNF|FR|NFR)[-\s]?(?P<number>\d{1,4})$', re.IGNORECASE)
PRIORITY_TYPE_RE = re.compile(
    r'^(?P<priority>alta|media|baja|critica|crítica)\s+(?P<type>no funcional|funcional)$',
    re.IGNORECASE,
)
PRIORITY_RE = re.compile(r'^(alta|media|baja|critica|crítica)$', re.IGNORECASE)
TYPE_RE = re.compile(r'^(no funcional|funcional)$', re.IGNORECASE)
TABLE_DESCRIPTION_START_RE = re.compile(
    r'\b(permitir|registrar|mostrar|guiar|mantener|calcular|vincular|crear|proporcionar|asignar|agrupar|ejecutar|gestionar|generar|estructurar|identificar|notificar|incluir|validar|garantizar|el sistema|la plataforma|la aplicacion)\b',
    re.IGNORECASE,
)


def extract_text_from_pdf(pdf_file):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RequirementPdfImportError(
            'No esta instalada la dependencia pypdf. Ejecuta pip install -r requirements/base.txt.'
        ) from exc

    try:
        pdf_file.seek(0)
        reader = PdfReader(pdf_file)
        pages = [page.extract_text() or '' for page in reader.pages]
    except Exception as exc:
        raise RequirementPdfImportError('No se pudo leer el PDF. Verifica que el archivo no este protegido o danado.') from exc

    text = '\n'.join(pages).strip()
    if not text:
        raise RequirementPdfImportError(
            'No se detecto texto en el PDF. Si es un documento escaneado, se requiere OCR.'
        )

    return text


def extract_text_from_source(source_file):
    """Extract plain text from a supported requirements document."""
    name = (getattr(source_file, 'name', '') or '').lower()
    extension = '.' + name.rsplit('.', 1)[-1] if '.' in name else ''
    try:
        source_file.seek(0)
        if extension == '.pdf':
            return extract_text_from_pdf(source_file)
        if extension in {'.txt', '.csv'}:
            raw = source_file.read()
            if isinstance(raw, str):
                return raw
            decoded = raw.decode('utf-8-sig', errors='replace')
            if extension == '.csv':
                rows = csv.reader(io.StringIO(decoded))
                return '\n'.join(' | '.join(cell.strip() for cell in row) for row in rows)
            return decoded.strip()
        if extension == '.docx':
            try:
                from docx import Document
            except ImportError as exc:
                raise RequirementSourceImportError('Instala python-docx para leer archivos Word (.docx).') from exc
            document = Document(source_file)
            return '\n'.join(p.text.strip() for p in document.paragraphs if p.text.strip())
        if extension in {'.xlsx', '.xls'}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RequirementSourceImportError('Instala openpyxl para leer archivos Excel.') from exc
            if extension == '.xls':
                raise RequirementSourceImportError('El formato .xls requiere guardarse como .xlsx antes de importarlo.')
            workbook = load_workbook(source_file, read_only=True, data_only=True)
            lines = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        lines.append(' | '.join(values))
            return '\n'.join(lines).strip()
        if extension == '.doc':
            if not shutil.which('libreoffice'):
                raise RequirementSourceImportError('El formato .doc requiere LibreOffice o convertirse a .docx.')
            with tempfile.TemporaryDirectory() as directory:
                input_path = f'{directory}/requirements.doc'
                with open(input_path, 'wb') as handle:
                    handle.write(source_file.read())
                result = subprocess.run(
                    ['libreoffice', '--headless', '--convert-to', 'txt:Text', '--outdir', directory, input_path],
                    capture_output=True,
                    timeout=30,
                    check=False,
                )
                output_path = f'{directory}/requirements.txt'
                if result.returncode != 0 or not os.path.exists(output_path):
                    raise RequirementSourceImportError('No se pudo leer el archivo Word .doc.')
                with open(output_path, encoding='utf-8', errors='replace') as handle:
                    return handle.read().strip()
    except RequirementSourceImportError:
        raise
    except Exception as exc:
        raise RequirementSourceImportError('No se pudo leer el documento. Verifica que no esté protegido o dañado.') from exc

    raise RequirementSourceImportError('Formato no soportado. Usa PDF, Word, Excel, CSV o TXT.')


def parse_requirements_from_text(text, defaults=None, limit=100):
    defaults = defaults or {}
    table_items = _parse_table_requirements(text, defaults, limit)
    if table_items:
        return table_items

    blocks = _split_requirement_blocks(text)
    parsed = []

    for block in blocks:
        item = _parse_requirement_block(block, defaults)
        if item:
            parsed.append(item)
        if len(parsed) >= limit:
            break

    return parsed


def _parse_table_requirements(text, defaults, limit):
    lines = [_clean_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    parsed = []
    index = 0

    while index < len(lines):
        code_info = _read_code_at(lines, index)
        if not code_info:
            index += 1
            continue

        code, index = code_info
        row_lines = []
        priority = defaults.get('priority')
        requirement_type = defaults.get('requirement_type')

        while index < len(lines) and not _read_code_at(lines, index):
            line = lines[index]
            priority_type = PRIORITY_TYPE_RE.match(line)
            if priority_type:
                priority = _priority_from_label(priority_type.group('priority')) or priority
                requirement_type = _type_from_label(priority_type.group('type')) or requirement_type
                index += 1
                break

            if (
                PRIORITY_RE.match(line)
                and index + 1 < len(lines)
                and TYPE_RE.match(lines[index + 1])
            ):
                priority = _priority_from_label(line) or priority
                requirement_type = _type_from_label(lines[index + 1]) or requirement_type
                index += 2
                break

            if not _is_table_header_line(line):
                row_lines.append(line)
            index += 1

        item = _build_table_requirement(code, row_lines, defaults, priority, requirement_type)
        if item:
            parsed.append(item)
        if len(parsed) >= limit:
            break

    # A single detected code can be a normal paragraph. Use table parsing only
    # when there is enough structure to prove the PDF extraction came from rows.
    return parsed if len(parsed) > 1 else []


def _read_code_at(lines, index):
    line = lines[index]
    complete = COMPLETE_CODE_RE.match(line)
    if complete:
        return _format_code(complete.group('prefix'), complete.group('number')), index + 1

    prefix = CODE_PREFIX_RE.match(line)
    if prefix and index + 1 < len(lines) and CODE_NUMBER_RE.match(lines[index + 1]):
        return _format_code(prefix.group('prefix'), lines[index + 1]), index + 2

    return None


def _format_code(prefix, number):
    return f'{prefix.upper()}-{number.zfill(3)}'


def _is_table_header_line(line):
    lowered = line.lower()
    return lowered in {
        'codigo',
        'código',
        'nombre del',
        'requisito',
        'descripción',
        'prioridad',
        'tipo',
    }


def _build_table_requirement(code, row_lines, defaults, priority, requirement_type):
    title, description = _split_table_title_description(row_lines)
    if not title or not description:
        return None

    return ParsedRequirement(
        title=title[:180],
        description=description,
        requirement_type=requirement_type or _detect_requirement_type(description, defaults.get('requirement_type'), code=code),
        priority=priority or _detect_priority(description, defaults.get('priority')),
        status=defaults.get('status', Requirement.Status.PENDING),
    )


def _split_table_title_description(row_lines):
    for index, line in enumerate(row_lines):
        match = TABLE_DESCRIPTION_START_RE.search(line)
        if not match:
            continue

        title_parts = row_lines[:index]
        description_parts = row_lines[index:]
        if match.start() > 0:
            title_parts.append(line[: match.start()].strip())
            description_parts = [line[match.start() :].strip(), *row_lines[index + 1 :]]

        title = _join_parts(title_parts)
        description = _join_parts(description_parts)
        return title, description

    description = _join_parts(row_lines)
    return _build_title(description), description


def _join_parts(parts):
    return re.sub(r'\s+', ' ', ' '.join(part for part in parts if part)).strip(' .:-')


def _priority_from_label(label):
    normalized = label.strip().lower()
    if normalized == 'alta':
        return Requirement.Priority.HIGH
    if normalized == 'media':
        return Requirement.Priority.MEDIUM
    if normalized == 'baja':
        return Requirement.Priority.LOW
    if normalized in {'critica', 'crítica'}:
        return Requirement.Priority.CRITICAL
    return None


def _type_from_label(label):
    normalized = label.strip().lower()
    if normalized == 'no funcional':
        return Requirement.RequirementType.NON_FUNCTIONAL
    if normalized == 'funcional':
        return Requirement.RequirementType.FUNCTIONAL
    return None


def _split_requirement_blocks(text):
    blocks = []
    current = []

    for raw_line in text.splitlines():
        line = _clean_line(raw_line)
        if not line:
            continue

        starts_requirement = _looks_like_requirement_start(line)
        if starts_requirement and current:
            blocks.append(' '.join(current))
            current = [line]
        elif starts_requirement:
            current = [line]
        elif current:
            current.append(line)

    if current:
        blocks.append(' '.join(current))

    return blocks


def _looks_like_requirement_start(line):
    if line.lower().startswith(('nota:', 'observacion:', 'comentario:')):
        return False

    match = REQUIREMENT_START_RE.match(line)
    if match and match.group('code'):
        return True

    return bool(NUMBERED_START_RE.match(line) and REQUIREMENT_WORD_RE.search(line)) or bool(
        REQUIREMENT_WORD_RE.search(line)
    )


def _parse_requirement_block(block, defaults):
    match = REQUIREMENT_START_RE.match(block)
    code = (match.group('code') or '') if match else ''
    body = match.group('body').strip() if match else block.strip()
    body = re.sub(r'^\d+[\).:-]\s+', '', body).strip()

    if body.lower().startswith(('nota:', 'observacion:', 'comentario:')):
        return None

    if len(body) < 12 or not REQUIREMENT_WORD_RE.search(body):
        return None

    title = _build_title(body)
    return ParsedRequirement(
        title=title,
        description=body,
        requirement_type=_detect_requirement_type(body, defaults.get('requirement_type'), code=code),
        priority=_detect_priority(body, defaults.get('priority')),
        status=defaults.get('status', Requirement.Status.PENDING),
    )


def _build_title(text):
    title = re.split(r'(?<=[.!?])\s+', text, maxsplit=1)[0]
    if ':' in title and len(title.split(':', 1)[0]) <= 90:
        title = title.split(':', 1)[0]

    title = title.strip(' .:-')
    if len(title) > 120:
        title = f'{title[:117].rstrip()}...'

    return title


def _detect_requirement_type(text, default, code=''):
    if code.upper().startswith(('RNF', 'NFR')):
        return Requirement.RequirementType.NON_FUNCTIONAL

    lowered = text.lower()
    non_functional_terms = (
        'no funcional',
        'rendimiento',
        'seguridad',
        'usabilidad',
        'disponibilidad',
        'escalabilidad',
        'accesibilidad',
        'compatibilidad',
        'mantenibilidad',
        'confiabilidad',
        'fiabilidad',
        'portabilidad',
        'latencia',
        'cifrar',
        'cifrado',
        'encriptar',
        'encriptacion',
        'autenticacion multifactor',
        'auditoria',
        'backup',
        'respaldo',
        'recuperacion',
        'integridad',
        'privacidad',
        'tiempo de respuesta',
        'concurrencia',
    )
    if any(term in lowered for term in non_functional_terms):
        return Requirement.RequirementType.NON_FUNCTIONAL

    return default or Requirement.RequirementType.FUNCTIONAL


def _detect_priority(text, default):
    lowered = text.lower()
    if 'critica' in lowered or 'critico' in lowered:
        return Requirement.Priority.CRITICAL
    if 'alta' in lowered:
        return Requirement.Priority.HIGH
    if 'baja' in lowered:
        return Requirement.Priority.LOW
    return default or Requirement.Priority.MEDIUM


def _clean_line(line):
    return re.sub(r'\s+', ' ', line).strip(' \t-')
